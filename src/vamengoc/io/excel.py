"""Type-preserving spreadsheet reader.

Workbooks are read cell by cell without letting pandas infer types: every cell
keeps the Python value the file stores (``str``, ``int``, ``float``,
``datetime``) and conversion happens later in explicit, logged parsers. The
reader never writes to its input (raw data are treated as immutable).

The Rust-based ``python-calamine`` reader is used when available (an order of
magnitude faster on the ~70 000-row AEFI series); ``openpyxl``/``xlrd`` remain
as fallbacks and give identical values after parsing (verified in the tests).
"""

from __future__ import annotations

import csv
import datetime as dt
from dataclasses import dataclass
from pathlib import Path
from typing import Any

__all__ = ["SheetGrid", "list_sheets", "read_sheet"]

_EXCEL_SUFFIXES = {".xlsx", ".xlsm"}


@dataclass(frozen=True)
class SheetGrid:
    """A rectangular grid of raw cell values (row-major, 0-based)."""

    path: Path
    sheet: str
    rows: list[list[Any]]

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def cell(self, r: int, c: int) -> Any:
        row = self.rows[r] if 0 <= r < len(self.rows) else []
        return row[c] if 0 <= c < len(row) else None


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _trim(rows: list[list[Any]]) -> list[list[Any]]:
    """Drop trailing empty rows and columns (Excel 'used range' artefacts)."""
    while rows and all(_is_blank(v) for v in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for j in range(len(r) - 1, -1, -1):
            if not _is_blank(r[j]):
                width = max(width, j + 1)
                break
    return [list(r[:width]) + [None] * (width - len(r[:width])) for r in rows]


def _calamine_rows(path: Path, sheet: str | None) -> tuple[str, list[list[Any]]] | None:
    try:
        from python_calamine import CalamineWorkbook
    except ImportError:  # pragma: no cover - optional accelerator
        return None
    wb = CalamineWorkbook.from_path(str(path))
    name = sheet if sheet is not None else wb.sheet_names[0]
    rows = wb.get_sheet_by_name(name).to_python(skip_empty_area=False)
    return name, [[None if (isinstance(v, str) and v == "") else v for v in r] for r in rows]


def list_sheets(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix in _EXCEL_SUFFIXES:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return list(book.sheet_names())
        finally:
            book.release_resources()
    if suffix == ".csv":
        return [path.stem]
    msg = f"Unsupported spreadsheet format: {path.name}"
    raise ValueError(msg)


def read_sheet(path: Path, sheet: str | None = None, *, engine: str = "auto") -> SheetGrid:
    """Read one sheet (default: the first) as a raw grid of cell values.

    ``engine`` is ``"auto"`` (calamine, else openpyxl/xlrd), ``"calamine"`` or
    ``"openpyxl"``.
    """
    suffix = path.suffix.lower()
    if engine in ("auto", "calamine") and suffix in {*_EXCEL_SUFFIXES, ".xls"}:
        got = _calamine_rows(path, sheet)
        if got is not None:
            return SheetGrid(path, got[0], _trim(got[1]))
        if engine == "calamine":  # pragma: no cover - explicit request without the package
            msg = "python-calamine is not installed"
            raise RuntimeError(msg)
    if suffix in _EXCEL_SUFFIXES:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            name = sheet if sheet is not None else wb.sheetnames[0]
            ws = wb[name]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        return SheetGrid(path, name, _trim(rows))
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        name = sheet if sheet is not None else book.sheet_names()[0]
        ws = book.sheet_by_name(name)
        out: list[list[Any]] = []
        for r in range(ws.nrows):
            row: list[Any] = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    v = float(cell.value)
                    row.append(int(v) if v.is_integer() else v)
                else:
                    row.append(cell.value)
            out.append(row)
        return SheetGrid(path, name, _trim(out))
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            rows_csv: list[list[Any]] = [[v if v != "" else None for v in r] for r in csv.reader(fh)]
        return SheetGrid(path, path.stem, _trim(rows_csv))
    msg = f"Unsupported spreadsheet format: {path.name}"
    raise ValueError(msg)


def is_datetime(v: Any) -> bool:
    return isinstance(v, dt.datetime | dt.date)
