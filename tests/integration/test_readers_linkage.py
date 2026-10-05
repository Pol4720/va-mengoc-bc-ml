"""Spreadsheet readers, normalisation dictionaries and lot linkage."""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook

from vamengoc.config import Project
from vamengoc.curate.linkage import link_reports_to_lots
from vamengoc.io.excel import list_sheets, read_sheet
from vamengoc.normalize.dictionaries import Mapper, split_multi
from vamengoc.normalize.registry import build_normalizers
from vamengoc.parsing.dates import parse_date
from vamengoc.parsing.text import normalize_text
from vamengoc.provenance import RunContext
from vamengoc.synth import SyntheticDataset


def test_engines_agree(synthetic_dir: SyntheticDataset) -> None:
    path = synthetic_dir.directory / "EA_2017_TOTAL.xlsx"
    a = read_sheet(path, "TOTAL", engine="calamine")
    b = read_sheet(path, "TOTAL", engine="openpyxl")
    assert a.n_rows == b.n_rows and a.n_cols == b.n_cols
    for ra, rb in zip(a.rows[:200], b.rows[:200], strict=True):
        for va, vb in zip(ra, rb, strict=True):
            da, db = parse_date(va), parse_date(vb)
            if isinstance(va, dt.date | dt.datetime) or isinstance(vb, dt.date | dt.datetime):
                assert da.value == db.value
            else:
                assert normalize_text(va) == normalize_text(vb)


def test_csv_and_xls_paths(tmp_path: Path) -> None:
    csv = tmp_path / "EA_2019_TOTAL.csv"
    csv.write_text("SEXO,EDAD\nM,6M\n,\n", encoding="utf-8")
    g = read_sheet(csv)
    assert list_sheets(csv) == ["EA_2019_TOTAL"] and g.rows[1] == ["M", "6M"] and g.n_rows == 2
    assert g.cell(10, 10) is None
    with pytest.raises(ValueError, match="Unsupported"):
        read_sheet(tmp_path / "x.docx")
    with pytest.raises(ValueError, match="Unsupported"):
        list_sheets(tmp_path / "x.docx")
    wb = Workbook()
    ws = wb.active
    ws.append(["a", None, None])
    ws.append([1, 2, None])
    ws.append([None, None, None])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    g2 = read_sheet(p, engine="openpyxl")
    assert g2.n_rows == 2 and g2.n_cols == 2 and list_sheets(p) == ["Sheet"]


def test_mapper_rules_and_log() -> None:
    m = Mapper(
        {"FINLAY": {"patterns": [".*FINLAY.*"], "fuzzy_anchors": ["FINLAY"]}, "UNKNOWN": {"patterns": ["", "N"]}},
        match_on="compact",
        fuzzy_threshold=80,
        default="OTHER",
    )
    assert m.map("Instituto Finlay") == "FINLAY"
    assert m.map("Finlai") == "FINLAY"  # fuzzy
    assert m.map("N") == "UNKNOWN"
    assert m.map(None) == "UNKNOWN"
    assert m.map("Serum") == "OTHER"
    assert m.map("Instituto Finlay") == "FINLAY"  # cached
    summ = m.log.summary()
    assert (
        summ["n_values"] == 6 and "pattern" in summ["by_rule"] and any(r.startswith("fuzzy") for r in summ["by_rule"])
    )
    assert m.log.table()[0]["n"] >= 1
    m2 = Mapper({"A": {"patterns": ["A"]}})
    assert m2.map("") is None and m2.map("zzz") is None
    assert split_multi("PENTA-L, IPV", r"\s*,\s*") == ["PENTA-L", "IPV"]
    assert split_multi(None, ",") == [] and split_multi("-", ",") == []


def test_normalizers_cover_documented_variants() -> None:
    norm = build_normalizers(Project())
    codes, unmapped = norm.vaccines("PENTA-L,IPV")
    assert codes == ["PENTA-L", "IPV"] and not unmapped
    assert norm.vaccines("Hib+DPT")[0] == ["HIB", "DPT"]
    assert norm.vaccines("am-bc")[0] == ["AM-BC"]
    assert norm.vaccines("VA-MENGOC-BC")[0] == ["AM-BC"]
    codes, unmapped = norm.vaccines("ZZTOP")
    assert codes == ["X:ZZTOP"] and unmapped == ["ZZTOP"]
    for raw in ("Heber Biotec", "heberbiotec", "heber- biotec", "CIGB"):
        assert norm.manufacturers(raw) == ["CIGB"]
    for raw in ("Finlay", "Instituto Finlay", "IFV", "Intituto Finlay"):
        assert norm.manufacturers(raw) == ["FINLAY"]
    assert norm.manufacturers("Serum Institute of India") == ["SII"]
    assert norm.manufacturers("N") == ["UNKNOWN"]
    assert norm.province.map("lh") == "LH" and norm.province.map("Camagüey") == "CM"
    assert norm.province_region["GT"] == "Oriente"
    assert norm.sites("1/3 CALM, BOCA") == ["THIGH", "ORAL"]
    assert norm.routes("IM,ORAL") == ["IM", "ORAL"]
    assert norm.place.map("Vacunatorio ") == "VACCINATION_ROOM"
    assert sorted(norm.destinations("Nacional-Colombia-Guatemala")) == ["CO", "CU", "GT"]
    assert norm.destinations(" Vietnam") == ["VN"]


def _reports(rows: list[dict]) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    df["has_target"] = True
    df["in_study_window"] = True
    for c in ("age_months", "ev_fever_39", "hospitalized", "coadministered", "sex", "target_lot_assignment"):
        if c not in df:
            df[c] = {
                "age_months": 3.0,
                "ev_fever_39": True,
                "hospitalized": False,
                "coadministered": False,
                "sex": "F",
                "target_lot_assignment": "single_vaccine",
            }[c]
    return df


def test_linkage_levels(tmp_path: Path) -> None:
    project = Project(
        overrides={"paths": {"runs_dir": str(tmp_path)}, "analysis": {"linkage": {"fuzzy_threshold": 85}}}
    )
    ctx = RunContext(project, "test", synthetic=True)
    lots = pd.DataFrame(
        {
            "lot_key": ["575M", "576M", "M101@2011#1", "M101@2016#1", "900M", "901M"],
            "lot_exact": ["575M", "576M", "M101", "M101", "900M", "901M"],
            "lot_core": ["575:M", "576:M", "101:M", "101:M", "900:M", "901:M"],
            "production_year": [2016, 2016, 2011, 2016, 2016, 2016],
            "national": [True, True, True, True, False, False],
        }
    )
    reps = _reports(
        [
            {"record_id": "r1", "analytic_year": 2017, "target_lots": "575M", "target_lots_core": "575:M"},
            {"record_id": "r2", "analytic_year": 2017, "target_lots": "M575", "target_lots_core": "575:M"},
            {"record_id": "r3", "analytic_year": 2017, "target_lots": "M101", "target_lots_core": "101:M"},
            {"record_id": "r4", "analytic_year": 2017, "target_lots": "5750M", "target_lots_core": "5750:M"},
            {"record_id": "r5", "analytic_year": 2017, "target_lots": None, "target_lots_core": None},
            {"record_id": "r6", "analytic_year": 2017, "target_lots": "575M|576M", "target_lots_core": "575:M|576:M"},
            {"record_id": "r7", "analytic_year": 2025, "target_lots": "575M", "target_lots_core": "575:M"},
            {"record_id": "r8", "analytic_year": 2017, "target_lots": "ZZ9", "target_lots_core": "9:ZZ"},
        ]
    )
    res = link_reports_to_lots(reps, lots, project, ctx)
    got = res.links.set_index("record_id")
    assert got.loc["r1", "match_level"] == "exact" and got.loc["r1", "lot_key"] == "575M"
    assert got.loc["r2", "match_level"] == "core" and got.loc["r2", "lot_key"] == "575M"
    assert got.loc["r3", "lot_key"] == "M101@2016#1"  # national + most recent <= vaccination year
    assert got.loc["r4", "match_level"] == "fuzzy"
    assert got.loc["r5", "match_level"] == "no_lot_recorded"
    assert got.loc["r6", "match_level"] == "multiple_lots"
    assert got.loc["r7", "temporal_ok"] == False  # noqa: E712 - produced 2016, used 2025
    assert got.loc["r8", "match_level"] == "no_match"
    assert res.summary["n_linked"] == 5 and res.summary["temporal_implausible"] == 1
    assert set(res.bias_table["linked"]) == {True, False}


def _lots_copy(synthetic_dir: SyntheticDataset, tmp_path: Path, rename: dict[str, str]) -> Path:
    """Copy of the synthetic lot-release workbook with some lot-sheet headers renamed."""
    from openpyxl import load_workbook

    src = next(synthetic_dir.directory.glob("BD*.xlsx"))
    wb = load_workbook(src)
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=6):
            for cell in row:
                for old, new in rename.items():
                    if isinstance(cell.value, str) and old.lower() in cell.value.lower():
                        cell.value = new
    out = tmp_path / src.name
    wb.save(out)
    return out


def test_lots_workbook_without_a_recognisable_attribute(synthetic_dir: SyntheticDataset, tmp_path: Path) -> None:
    from vamengoc.ingest.lots import ingest_lots_workbook

    project = Project(overrides={"paths": {"runs_dir": str(tmp_path)}})
    path = _lots_copy(synthetic_dir, tmp_path, {"act. antic. IgG": "INMUNOG. XYZ (ANTICUERPOS)"})
    ctx = RunContext(project, "test", synthetic=True)
    res = ingest_lots_workbook(project, path, build_normalizers(project), ctx)
    assert "igg_elisa" not in res.lots.columns
    assert {"endotoxin", "bactericidal_titer"} <= set(res.lots.columns)
    assert any(s["step"] == "lots_attributes_not_found" and s["attributes"] == ["igg_elisa"] for s in ctx.steps)


def test_lots_workbook_with_too_few_attributes_stops(synthetic_dir: SyntheticDataset, tmp_path: Path) -> None:
    from vamengoc.ingest.lots import ingest_lots_workbook
    from vamengoc.io.schema_registry import SchemaError

    project = Project(overrides={"paths": {"runs_dir": str(tmp_path)}})
    rename = {
        k: f"COLUMNA {i}" for i, k in enumerate(["Inmunogenicidad", "endotoxinas", "tiomersal", "bulbo", "Al(OH)"])
    }
    path = _lots_copy(synthetic_dir, tmp_path, rename)
    with pytest.raises(SchemaError, match="quality attributes recognised"):
        ingest_lots_workbook(project, path, build_normalizers(project), RunContext(project, "test", synthetic=True))
